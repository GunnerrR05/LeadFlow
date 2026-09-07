import os
from datetime import date, datetime
from shutil import copy2

from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from config import (
    ARQUIVO_LEADS,
    PASTA_BACKUPS,
    preparar_ambiente,
)


CABECALHOS = [
    "DATA DE CADASTRO",
    "PRÓXIMO CONTATO",
    "ÚLTIMA INTERAÇÃO",
    "UNIDADE",
    "STATUS",
    "TELEFONE",
    "NOME",
    "PRODUTO",
    "CIDADE / UF",
    "E-MAIL",
    "APLICAÇÃO",
    "ORIGEM",
    "CONSULTOR",
    "OBSERVAÇÃO",
    "RESUMO",
    "NOME NO WHATSAPP",
]


STATUS_VALIDOS = [
    "EM ANDAMENTO",
    "NEGOCIAÇÃO",
    "DECLINADO",
    "CONQUISTADO - SUPRIM",
    "CONQUISTADO - EQUIP",
    "FUTURA",
    "CLIENTE ATIVO",
]


LARGURAS_COLUNAS = {
    "A": 20,
    "B": 18,
    "C": 20,
    "D": 20,
    "E": 24,
    "F": 18,
    "G": 24,
    "H": 24,
    "I": 18,
    "J": 28,
    "K": 24,
    "L": 18,
    "M": 18,
    "N": 45,
    "O": 55,
    "P": 65,
}


preparar_ambiente()


def _valor(campo):
    """
    Converte valores para texto.

    Também aceita objetos Entry e Combobox
    do Tkinter.
    """

    if (
        hasattr(campo, "get")
        and callable(campo.get)
    ):
        campo = campo.get()

    if campo is None:
        return ""

    if isinstance(campo, datetime):
        if (
            campo.hour == 0
            and campo.minute == 0
            and campo.second == 0
        ):
            return campo.strftime(
                "%d/%m/%Y"
            )

        return campo.strftime(
            "%d/%m/%Y %H:%M"
        )

    if isinstance(campo, date):
        return campo.strftime(
            "%d/%m/%Y"
        )

    return str(campo).strip()


def _normalizar_telefone(valor):
    """
    Mantém somente os números do telefone.
    """

    return "".join(
        caractere
        for caractere in _valor(valor)
        if caractere.isdigit()
    )


def _configurar_planilha(planilha):
    """
    Aplica formatação e configurações
    na planilha principal.
    """

    planilha.freeze_panes = "A2"

    ultima_linha = max(
        planilha.max_row,
        1
    )

    planilha.auto_filter.ref = (
        f"A1:P{ultima_linha}"
    )

    for coluna, largura in (
        LARGURAS_COLUNAS.items()
    ):
        planilha.column_dimensions[
            coluna
        ].width = largura

    preenchimento_cabecalho = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    for celula in planilha[1]:
        celula.font = Font(
            bold=True
        )

        celula.fill = (
            preenchimento_cabecalho
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    planilha.row_dimensions[
        1
    ].height = 30

    if planilha.max_row >= 2:
        for linha in planilha.iter_rows(
            min_row=2,
            max_row=planilha.max_row,
            max_col=16
        ):
            for celula in linha:
                celula.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

    # Telefone agora está na coluna F.
    for linha in range(
        2,
        planilha.max_row + 1
    ):
        planilha[
            f"F{linha}"
        ].number_format = "@"


def _salvar_seguro(arquivo):
    """
    Salva primeiro em um arquivo temporário
    e depois substitui a planilha principal.
    """

    preparar_ambiente()

    arquivo_temporario = (
        ARQUIVO_LEADS.with_name(
            "leads_salvando.tmp.xlsx"
        )
    )

    try:
        if arquivo_temporario.exists():
            arquivo_temporario.unlink()

        arquivo.save(
            arquivo_temporario
        )

        arquivo.close()

        os.replace(
            arquivo_temporario,
            ARQUIVO_LEADS
        )

    except PermissionError as erro:
        raise PermissionError(
            "Não foi possível salvar os dados.\n\n"
            "Feche o arquivo leads.xlsx no Excel "
            "e tente novamente."
        ) from erro

    finally:
        try:
            arquivo.close()

        except Exception:
            pass

        if arquivo_temporario.exists():
            try:
                arquivo_temporario.unlink()

            except OSError:
                pass


def _criar_backup_migracao():
    """
    Cria um backup antes de acrescentar
    a coluna UNIDADE.
    """

    preparar_ambiente()

    PASTA_BACKUPS.mkdir(
        parents=True,
        exist_ok=True
    )

    data_hora = datetime.now().strftime(
        "%Y-%m-%d_%H-%M-%S-%f"
    )

    destino = (
        PASTA_BACKUPS
        / (
            "leads_backup_antes_unidade_"
            f"{data_hora}.xlsx"
        )
    )

    copy2(
        ARQUIVO_LEADS,
        destino
    )

    return destino


def _migrar_estrutura(
    arquivo,
    planilha
):
    """
    Migra a planilha antiga:

    Antes:
    D = STATUS
    E = TELEFONE

    Depois:
    D = UNIDADE
    E = STATUS
    F = TELEFONE
    """

    cabecalho_d = _valor(
        planilha["D1"].value
    ).upper()

    cabecalho_e = _valor(
        planilha["E1"].value
    ).upper()

    # A planilha já está atualizada.
    if (
        cabecalho_d == "UNIDADE"
        and cabecalho_e == "STATUS"
    ):
        return (
            arquivo,
            planilha
        )

    # Estrutura antiga conhecida.
    if cabecalho_d == "STATUS":
        try:
            _criar_backup_migracao()

            # Insere uma nova coluna antes
            # da antiga coluna STATUS.
            planilha.insert_cols(
                4,
                1
            )

            for coluna, cabecalho in enumerate(
                CABECALHOS,
                start=1
            ):
                planilha.cell(
                    row=1,
                    column=coluna,
                    value=cabecalho
                )

            _configurar_planilha(
                planilha
            )

            _salvar_seguro(
                arquivo
            )

        except Exception:
            try:
                arquivo.close()

            except Exception:
                pass

            raise

        # Reabre o arquivo após a migração.
        arquivo_novo = load_workbook(
            ARQUIVO_LEADS,
            data_only=False
        )

        return (
            arquivo_novo,
            arquivo_novo.active
        )

    arquivo.close()

    raise RuntimeError(
        "A estrutura da planilha não foi reconhecida.\n\n"
        "Nenhuma alteração foi realizada.\n\n"
        "Restaure um backup ou verifique os "
        "cabeçalhos da planilha."
    )


def _criar_planilha_nova():
    arquivo = Workbook()
    planilha = arquivo.active
    planilha.title = "Leads"

    for coluna, cabecalho in enumerate(
        CABECALHOS,
        start=1
    ):
        planilha.cell(
            row=1,
            column=coluna,
            value=cabecalho
        )

    _configurar_planilha(
        planilha
    )

    _salvar_seguro(
        arquivo
    )


def _abrir_planilha():
    """
    Abre a planilha e realiza a migração
    automaticamente quando necessário.
    """

    preparar_ambiente()

    if not ARQUIVO_LEADS.exists():
        _criar_planilha_nova()

    try:
        arquivo = load_workbook(
            ARQUIVO_LEADS,
            data_only=False
        )

    except PermissionError as erro:
        raise PermissionError(
            "Não foi possível abrir leads.xlsx.\n\n"
            "Feche a planilha no Excel "
            "e tente novamente."
        ) from erro

    except Exception as erro:
        raise RuntimeError(
            "Não foi possível abrir a planilha "
            "de leads.\n\n"
            f"Arquivo: {ARQUIVO_LEADS}"
        ) from erro

    planilha = arquivo.active

    return _migrar_estrutura(
        arquivo,
        planilha
    )


def _gerar_resumo(dados):
    """
    Formato antigo do resumo.
    """

    return (
        f"{dados['nome']} - "
        f"{dados['telefone']} - "
        f"{dados['origem']} - "
        f"{dados['produto']}"
    )


def _gerar_nome_whatsapp(dados):
    """
    Formato do nome para o WhatsApp.
    """

    return (
        "LEAD - "
        f"{dados['nome']} - "
        f"{dados['produto']} - "
        f"{dados['cidade_uf']} - "
        f"{dados['origem']}"
    )


def _lead_da_linha(
    planilha,
    linha
):
    """
    Transforma uma linha do Excel
    em um dicionário de lead.
    """

    return {
        "linha": linha,

        "data_cadastro": _valor(
            planilha[f"A{linha}"].value
        ),

        "proximo_contato": _valor(
            planilha[f"B{linha}"].value
        ),

        "ultima_interacao": _valor(
            planilha[f"C{linha}"].value
        ),

        "unidade": _valor(
            planilha[f"D{linha}"].value
        ),

        "status": (
            _valor(
                planilha[f"E{linha}"].value
            )
            or "EM ANDAMENTO"
        ),

        "telefone": _valor(
            planilha[f"F{linha}"].value
        ),

        "nome": _valor(
            planilha[f"G{linha}"].value
        ),

        "produto": _valor(
            planilha[f"H{linha}"].value
        ),

        "cidade_uf": _valor(
            planilha[f"I{linha}"].value
        ),

        "email": _valor(
            planilha[f"J{linha}"].value
        ),

        "aplicacao": _valor(
            planilha[f"K{linha}"].value
        ),

        "origem": _valor(
            planilha[f"L{linha}"].value
        ),

        "consultor": _valor(
            planilha[f"M{linha}"].value
        ),

        "observacao": _valor(
            planilha[f"N{linha}"].value
        ),

        "resumo": _valor(
            planilha[f"O{linha}"].value
        ),

        "whatsapp": _valor(
            planilha[f"P{linha}"].value
        ),
    }


def _montar_dados(
    lead,
    atual=None
):
    """
    Organiza os dados para cadastro
    ou atualização.
    """

    atual = atual or {}

    def obter(
        chave,
        padrao=""
    ):
        return _valor(
            lead.get(
                chave,
                atual.get(
                    chave,
                    padrao
                )
            )
        )

    dados = {
        "proximo_contato": obter(
            "proximo_contato"
        ),

        "ultima_interacao": obter(
            "ultima_interacao"
        ),

        "unidade": obter(
            "unidade"
        ),

        "status": (
            obter(
                "status",
                "EM ANDAMENTO"
            )
            or "EM ANDAMENTO"
        ),

        "telefone": obter(
            "telefone"
        ),

        "nome": obter(
            "nome"
        ),

        "produto": _valor(
            lead.get(
                "produto",
                lead.get(
                    "interesse",
                    atual.get(
                        "produto",
                        ""
                    )
                )
            )
        ),

        "cidade_uf": obter(
            "cidade_uf"
        ),

        "email": obter(
            "email"
        ),

        "aplicacao": obter(
            "aplicacao"
        ),

        "origem": obter(
            "origem"
        ),

        "consultor": obter(
            "consultor"
        ),

        "observacao": obter(
            "observacao"
        ),
    }

    # Regra central de segurança:
    # declinado não possui próximo contato.
    if (
        dados["status"]
        .strip()
        .upper()
        == "DECLINADO"
    ):
        dados[
            "proximo_contato"
        ] = ""

    return dados


def _escrever_linha(
    planilha,
    linha,
    data_cadastro,
    dados
):
    resumo = _gerar_resumo(
        dados
    )

    nome_whatsapp = (
        _gerar_nome_whatsapp(
            dados
        )
    )

    valores = [
        data_cadastro,
        dados["proximo_contato"],
        dados["ultima_interacao"],
        dados["unidade"],
        dados["status"],
        dados["telefone"],
        dados["nome"],
        dados["produto"],
        dados["cidade_uf"],
        dados["email"],
        dados["aplicacao"],
        dados["origem"],
        dados["consultor"],
        dados["observacao"],
        resumo,
        nome_whatsapp,
    ]

    for coluna, valor in enumerate(
        valores,
        start=1
    ):
        planilha.cell(
            row=linha,
            column=coluna,
            value=valor
        )

    planilha[
        f"F{linha}"
    ].number_format = "@"

    return (
        resumo,
        nome_whatsapp
    )


def criar_planilha():
    """
    Cria a planilha quando não existe.

    Quando já existe, verifica e realiza
    a migração da coluna UNIDADE.
    """

    preparar_ambiente()

    if ARQUIVO_LEADS.exists():
        arquivo, _ = (
            _abrir_planilha()
        )

        arquivo.close()

        return False

    _criar_planilha_nova()

    return True


def salvar_lead(lead):
    """
    Salva um novo lead.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        linha_vazia = (
            planilha.max_row + 1
        )

        dados = _montar_dados(
            lead
        )

        data_cadastro = (
            datetime.now().strftime(
                "%d/%m/%Y %H:%M"
            )
        )

        _escrever_linha(
            planilha,
            linha_vazia,
            data_cadastro,
            dados
        )

        _configurar_planilha(
            planilha
        )

        _salvar_seguro(
            arquivo
        )

        return linha_vazia

    except Exception:
        try:
            arquivo.close()

        except Exception:
            pass

        raise


def pegar_leads():
    """
    Retorna todos os leads cadastrados.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        leads = []

        for linha in range(
            2,
            planilha.max_row + 1
        ):
            telefone = (
                planilha[
                    f"F{linha}"
                ].value
            )

            nome = (
                planilha[
                    f"G{linha}"
                ].value
            )

            if (
                telefone is None
                and nome is None
            ):
                continue

            leads.append(
                _lead_da_linha(
                    planilha,
                    linha
                )
            )

        return leads

    finally:
        arquivo.close()


def listar_leads():
    """
    Mantida para compatibilidade.
    """

    return pegar_leads()


def buscar_lead_por_nome(
    nome_busca
):
    """
    Busca um lead pelo nome e retorna
    a linha da planilha.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        nome_procurado = (
            _valor(
                nome_busca
            ).casefold()
        )

        for linha in range(
            2,
            planilha.max_row + 1
        ):
            nome = _valor(
                planilha[
                    f"G{linha}"
                ].value
            )

            if (
                nome.casefold()
                == nome_procurado
            ):
                return linha

        return None

    finally:
        arquivo.close()


def buscar_lead_por_telefone(
    telefone_busca
):
    """
    Busca um lead pelo telefone,
    ignorando a formatação.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        telefone_procurado = (
            _normalizar_telefone(
                telefone_busca
            )
        )

        for linha in range(
            2,
            planilha.max_row + 1
        ):
            telefone_planilha = (
                _normalizar_telefone(
                    planilha[
                        f"F{linha}"
                    ].value
                )
            )

            if (
                telefone_planilha
                == telefone_procurado
            ):
                return _lead_da_linha(
                    planilha,
                    linha
                )

        return None

    finally:
        arquivo.close()


def atualizar_lead(lead):
    """
    Atualiza um lead pela chave linha.
    """

    if "linha" not in lead:
        raise KeyError(
            "O lead precisa possuir "
            "a chave 'linha'."
        )

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        linha = int(
            lead["linha"]
        )

        if (
            linha < 2
            or linha > planilha.max_row
        ):
            raise ValueError(
                "Linha do lead inválida."
            )

        atual = _lead_da_linha(
            planilha,
            linha
        )

        dados = _montar_dados(
            lead,
            atual
        )

        data_cadastro = (
            _valor(
                lead.get(
                    "data_cadastro",
                    lead.get(
                        "data",
                        atual[
                            "data_cadastro"
                        ]
                    )
                )
            )
            or atual[
                "data_cadastro"
            ]
        )

        resumo, whatsapp = (
            _escrever_linha(
                planilha,
                linha,
                data_cadastro,
                dados
            )
        )

        _configurar_planilha(
            planilha
        )

        _salvar_seguro(
            arquivo
        )

        dados[
            "data_cadastro"
        ] = data_cadastro

        dados["linha"] = linha
        dados["resumo"] = resumo
        dados["whatsapp"] = whatsapp

        return dados

    except Exception:
        try:
            arquivo.close()

        except Exception:
            pass

        raise


def excluir_lead(linha):
    """
    Exclui uma linha da planilha.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        linha = int(
            linha
        )

        if (
            linha < 2
            or linha > planilha.max_row
        ):
            raise ValueError(
                "Linha do lead inválida."
            )

        planilha.delete_rows(
            linha
        )

        _configurar_planilha(
            planilha
        )

        _salvar_seguro(
            arquivo
        )

        return True

    except Exception:
        try:
            arquivo.close()

        except Exception:
            pass

        raise


def atualizar_status(
    telefone,
    novo_status
):
    """
    Atualiza o status pelo telefone
    e registra a última interação.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        telefone_procurado = (
            _normalizar_telefone(
                telefone
            )
        )

        status_novo = _valor(
            novo_status
        )

        for linha in range(
            2,
            planilha.max_row + 1
        ):
            telefone_planilha = (
                _normalizar_telefone(
                    planilha[
                        f"F{linha}"
                    ].value
                )
            )

            if (
                telefone_planilha
                != telefone_procurado
            ):
                continue

            planilha[
                f"E{linha}"
            ] = status_novo

            planilha[
                f"C{linha}"
            ] = (
                datetime.now().strftime(
                    "%d/%m/%Y %H:%M"
                )
            )

            if (
                status_novo
                .strip()
                .upper()
                == "DECLINADO"
            ):
                planilha[
                    f"B{linha}"
                ] = ""

            _configurar_planilha(
                planilha
            )

            _salvar_seguro(
                arquivo
            )

            return True

        arquivo.close()

        return False

    except Exception:
        try:
            arquivo.close()

        except Exception:
            pass

        raise


def dashboard():
    """
    Retorna as quantidades por status
    e a taxa de conversão.
    """

    leads = pegar_leads()

    contagem = {
        status: 0
        for status in STATUS_VALIDOS
    }

    for lead in leads:
        status = (
            _valor(
                lead.get(
                    "status"
                )
            )
            or "EM ANDAMENTO"
        )

        if status not in contagem:
            contagem[status] = 0

        contagem[status] += 1

    total = len(
        leads
    )

    conquistados = (
        contagem.get(
            "CONQUISTADO - SUPRIM",
            0
        )
        +
        contagem.get(
            "CONQUISTADO - EQUIP",
            0
        )
    )

    taxa_conversao = (
        conquistados / total * 100
        if total
        else 0
    )

    return {
        "total": total,
        "status": contagem,
        "taxa_conversao": taxa_conversao,
    }


def _converter_data_contato(valor):
    """
    Converte o próximo contato para data.
    """

    if isinstance(
        valor,
        datetime
    ):
        return valor.date()

    if isinstance(
        valor,
        date
    ):
        return valor

    texto = _valor(
        valor
    )

    if not texto:
        return None

    for formato in (
        "%d/%m/%Y",
        "%d/%m/%Y %H:%M",
    ):
        try:
            return datetime.strptime(
                texto,
                formato
            ).date()

        except ValueError:
            continue

    return None


def listar_followups():
    """
    Função mantida para compatibilidade.

    Leads declinados ficam em Finalizado.
    Leads ativos sem data não são exibidos.
    """

    hoje = datetime.now().date()

    grupos = {
        "atrasados": [],
        "hoje": [],
        "proximos": [],
        "finalizados": [],
        "invalidos": [],
    }

    for lead in pegar_leads():
        item = lead.copy()

        status = _valor(
            lead.get(
                "status",
                ""
            )
        ).upper()

        if status == "DECLINADO":
            grupos[
                "finalizados"
            ].append(
                item
            )
            continue

        texto_data = _valor(
            lead.get(
                "proximo_contato",
                ""
            )
        )

        # Leads ativos sem data ficam fora
        # do Follow-up.
        if not texto_data:
            continue

        data_contato = (
            _converter_data_contato(
                texto_data
            )
        )

        if data_contato is None:
            grupos[
                "invalidos"
            ].append(
                item
            )
            continue

        item[
            "_data_contato"
        ] = data_contato

        if data_contato < hoje:
            grupo = "atrasados"

        elif data_contato == hoje:
            grupo = "hoje"

        else:
            grupo = "proximos"

        grupos[
            grupo
        ].append(
            item
        )

    def ordenar_data(item):
        return (
            item["_data_contato"],
            _valor(
                item.get(
                    "nome",
                    ""
                )
            ).casefold()
        )

    for grupo in (
        "atrasados",
        "hoje",
        "proximos",
    ):
        grupos[grupo].sort(
            key=ordenar_data
        )

        for item in grupos[grupo]:
            item.pop(
                "_data_contato",
                None
            )

    for grupo in (
        "finalizados",
        "invalidos",
    ):
        grupos[grupo].sort(
            key=lambda item: _valor(
                item.get(
                    "nome",
                    ""
                )
            ).casefold()
        )

    return grupos


if __name__ == "__main__":
    criado = criar_planilha()

    if criado:
        print(
            "Planilha criada com sucesso."
        )

    else:
        print(
            "Planilha verificada com sucesso."
        )

    print(
        ARQUIVO_LEADS
        )
import os
from datetime import date, datetime
from shutil import copy2

from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from config import (
    ARQUIVO_LEADS,
    PASTA_BACKUPS,
    preparar_ambiente,
)


CABECALHOS = [
    "DATA DE CADASTRO",
    "PRÓXIMO CONTATO",
    "ÚLTIMA INTERAÇÃO",
    "UNIDADE",
    "STATUS",
    "TELEFONE",
    "NOME",
    "PRODUTO",
    "CIDADE / UF",
    "E-MAIL",
    "APLICAÇÃO",
    "ORIGEM",
    "CONSULTOR",
    "OBSERVAÇÃO",
    "RESUMO",
    "NOME NO WHATSAPP",
]


STATUS_VALIDOS = [
    "EM ANDAMENTO",
    "NEGOCIAÇÃO",
    "DECLINADO",
    "CONQUISTADO - SUPRIM",
    "CONQUISTADO - EQUIP",
    "FUTURA",
    "CLIENTE ATIVO",
]


LARGURAS_COLUNAS = {
    "A": 20,
    "B": 18,
    "C": 20,
    "D": 20,
    "E": 24,
    "F": 18,
    "G": 24,
    "H": 24,
    "I": 18,
    "J": 28,
    "K": 24,
    "L": 18,
    "M": 18,
    "N": 45,
    "O": 55,
    "P": 65,
}


preparar_ambiente()


def _valor(campo):
    """
    Converte valores para texto.

    Também aceita objetos Entry e Combobox
    do Tkinter.
    """

    if (
        hasattr(campo, "get")
        and callable(campo.get)
    ):
        campo = campo.get()

    if campo is None:
        return ""

    if isinstance(campo, datetime):
        if (
            campo.hour == 0
            and campo.minute == 0
            and campo.second == 0
        ):
            return campo.strftime(
                "%d/%m/%Y"
            )

        return campo.strftime(
            "%d/%m/%Y %H:%M"
        )

    if isinstance(campo, date):
        return campo.strftime(
            "%d/%m/%Y"
        )

    return str(campo).strip()


def _normalizar_telefone(valor):
    """
    Mantém somente os números do telefone.
    """

    return "".join(
        caractere
        for caractere in _valor(valor)
        if caractere.isdigit()
    )


def _configurar_planilha(planilha):
    """
    Aplica formatação e configurações
    na planilha principal.
    """

    planilha.freeze_panes = "A2"

    ultima_linha = max(
        planilha.max_row,
        1
    )

    planilha.auto_filter.ref = (
        f"A1:P{ultima_linha}"
    )

    for coluna, largura in (
        LARGURAS_COLUNAS.items()
    ):
        planilha.column_dimensions[
            coluna
        ].width = largura

    preenchimento_cabecalho = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    for celula in planilha[1]:
        celula.font = Font(
            bold=True
        )

        celula.fill = (
            preenchimento_cabecalho
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    planilha.row_dimensions[
        1
    ].height = 30

    if planilha.max_row >= 2:
        for linha in planilha.iter_rows(
            min_row=2,
            max_row=planilha.max_row,
            max_col=16
        ):
            for celula in linha:
                celula.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

    # Telefone agora está na coluna F.
    for linha in range(
        2,
        planilha.max_row + 1
    ):
        planilha[
            f"F{linha}"
        ].number_format = "@"


def _salvar_seguro(arquivo):
    """
    Salva primeiro em um arquivo temporário
    e depois substitui a planilha principal.
    """

    preparar_ambiente()

    arquivo_temporario = (
        ARQUIVO_LEADS.with_name(
            "leads_salvando.tmp.xlsx"
        )
    )

    try:
        if arquivo_temporario.exists():
            arquivo_temporario.unlink()

        arquivo.save(
            arquivo_temporario
        )

        arquivo.close()

        os.replace(
            arquivo_temporario,
            ARQUIVO_LEADS
        )

    except PermissionError as erro:
        raise PermissionError(
            "Não foi possível salvar os dados.\n\n"
            "Feche o arquivo leads.xlsx no Excel "
            "e tente novamente."
        ) from erro

    finally:
        try:
            arquivo.close()

        except Exception:
            pass

        if arquivo_temporario.exists():
            try:
                arquivo_temporario.unlink()

            except OSError:
                pass


def _criar_backup_migracao():
    """
    Cria um backup antes de acrescentar
    a coluna UNIDADE.
    """

    preparar_ambiente()

    PASTA_BACKUPS.mkdir(
        parents=True,
        exist_ok=True
    )

    data_hora = datetime.now().strftime(
        "%Y-%m-%d_%H-%M-%S-%f"
    )

    destino = (
        PASTA_BACKUPS
        / (
            "leads_backup_antes_unidade_"
            f"{data_hora}.xlsx"
        )
    )

    copy2(
        ARQUIVO_LEADS,
        destino
    )

    return destino


def _migrar_estrutura(
    arquivo,
    planilha
):
    """
    Migra a planilha antiga:

    Antes:
    D = STATUS
    E = TELEFONE

    Depois:
    D = UNIDADE
    E = STATUS
    F = TELEFONE
    """

    cabecalho_d = _valor(
        planilha["D1"].value
    ).upper()

    cabecalho_e = _valor(
        planilha["E1"].value
    ).upper()

    # A planilha já está atualizada.
    if (
        cabecalho_d == "UNIDADE"
        and cabecalho_e == "STATUS"
    ):
        return (
            arquivo,
            planilha
        )

    # Estrutura antiga conhecida.
    if cabecalho_d == "STATUS":
        try:
            _criar_backup_migracao()

            # Insere uma nova coluna antes
            # da antiga coluna STATUS.
            planilha.insert_cols(
                4,
                1
            )

            for coluna, cabecalho in enumerate(
                CABECALHOS,
                start=1
            ):
                planilha.cell(
                    row=1,
                    column=coluna,
                    value=cabecalho
                )

            _configurar_planilha(
                planilha
            )

            _salvar_seguro(
                arquivo
            )

        except Exception:
            try:
                arquivo.close()

            except Exception:
                pass

            raise

        # Reabre o arquivo após a migração.
        arquivo_novo = load_workbook(
            ARQUIVO_LEADS,
            data_only=False
        )

        return (
            arquivo_novo,
            arquivo_novo.active
        )

    arquivo.close()

    raise RuntimeError(
        "A estrutura da planilha não foi reconhecida.\n\n"
        "Nenhuma alteração foi realizada.\n\n"
        "Restaure um backup ou verifique os "
        "cabeçalhos da planilha."
    )


def _criar_planilha_nova():
    arquivo = Workbook()
    planilha = arquivo.active
    planilha.title = "Leads"

    for coluna, cabecalho in enumerate(
        CABECALHOS,
        start=1
    ):
        planilha.cell(
            row=1,
            column=coluna,
            value=cabecalho
        )

    _configurar_planilha(
        planilha
    )

    _salvar_seguro(
        arquivo
    )


def _abrir_planilha():
    """
    Abre a planilha e realiza a migração
    automaticamente quando necessário.
    """

    preparar_ambiente()

    if not ARQUIVO_LEADS.exists():
        _criar_planilha_nova()

    try:
        arquivo = load_workbook(
            ARQUIVO_LEADS,
            data_only=False
        )

    except PermissionError as erro:
        raise PermissionError(
            "Não foi possível abrir leads.xlsx.\n\n"
            "Feche a planilha no Excel "
            "e tente novamente."
        ) from erro

    except Exception as erro:
        raise RuntimeError(
            "Não foi possível abrir a planilha "
            "de leads.\n\n"
            f"Arquivo: {ARQUIVO_LEADS}"
        ) from erro

    planilha = arquivo.active

    return _migrar_estrutura(
        arquivo,
        planilha
    )


def _gerar_resumo(dados):
    """
    Formato antigo do resumo.
    """

    return (
        f"{dados['nome']} - "
        f"{dados['telefone']} - "
        f"{dados['origem']} - "
        f"{dados['produto']}"
    )


def _gerar_nome_whatsapp(dados):
    """
    Formato do nome para o WhatsApp.
    """

    return (
        "LEAD - "
        f"{dados['nome']} - "
        f"{dados['produto']} - "
        f"{dados['cidade_uf']} - "
        f"{dados['origem']}"
    )


def _lead_da_linha(
    planilha,
    linha
):
    """
    Transforma uma linha do Excel
    em um dicionário de lead.
    """

    return {
        "linha": linha,

        "data_cadastro": _valor(
            planilha[f"A{linha}"].value
        ),

        "proximo_contato": _valor(
            planilha[f"B{linha}"].value
        ),

        "ultima_interacao": _valor(
            planilha[f"C{linha}"].value
        ),

        "unidade": _valor(
            planilha[f"D{linha}"].value
        ),

        "status": (
            _valor(
                planilha[f"E{linha}"].value
            )
            or "EM ANDAMENTO"
        ),

        "telefone": _valor(
            planilha[f"F{linha}"].value
        ),

        "nome": _valor(
            planilha[f"G{linha}"].value
        ),

        "produto": _valor(
            planilha[f"H{linha}"].value
        ),

        "cidade_uf": _valor(
            planilha[f"I{linha}"].value
        ),

        "email": _valor(
            planilha[f"J{linha}"].value
        ),

        "aplicacao": _valor(
            planilha[f"K{linha}"].value
        ),

        "origem": _valor(
            planilha[f"L{linha}"].value
        ),

        "consultor": _valor(
            planilha[f"M{linha}"].value
        ),

        "observacao": _valor(
            planilha[f"N{linha}"].value
        ),

        "resumo": _valor(
            planilha[f"O{linha}"].value
        ),

        "whatsapp": _valor(
            planilha[f"P{linha}"].value
        ),
    }


def _montar_dados(
    lead,
    atual=None
):
    """
    Organiza os dados para cadastro
    ou atualização.
    """

    atual = atual or {}

    def obter(
        chave,
        padrao=""
    ):
        return _valor(
            lead.get(
                chave,
                atual.get(
                    chave,
                    padrao
                )
            )
        )

    dados = {
        "proximo_contato": obter(
            "proximo_contato"
        ),

        "ultima_interacao": obter(
            "ultima_interacao"
        ),

        "unidade": obter(
            "unidade"
        ),

        "status": (
            obter(
                "status",
                "EM ANDAMENTO"
            )
            or "EM ANDAMENTO"
        ),

        "telefone": obter(
            "telefone"
        ),

        "nome": obter(
            "nome"
        ),

        "produto": _valor(
            lead.get(
                "produto",
                lead.get(
                    "interesse",
                    atual.get(
                        "produto",
                        ""
                    )
                )
            )
        ),

        "cidade_uf": obter(
            "cidade_uf"
        ),

        "email": obter(
            "email"
        ),

        "aplicacao": obter(
            "aplicacao"
        ),

        "origem": obter(
            "origem"
        ),

        "consultor": obter(
            "consultor"
        ),

        "observacao": obter(
            "observacao"
        ),
    }

    # Regra central de segurança:
    # declinado não possui próximo contato.
    if (
        dados["status"]
        .strip()
        .upper()
        == "DECLINADO"
    ):
        dados[
            "proximo_contato"
        ] = ""

    return dados


def _escrever_linha(
    planilha,
    linha,
    data_cadastro,
    dados
):
    resumo = _gerar_resumo(
        dados
    )

    nome_whatsapp = (
        _gerar_nome_whatsapp(
            dados
        )
    )

    valores = [
        data_cadastro,
        dados["proximo_contato"],
        dados["ultima_interacao"],
        dados["unidade"],
        dados["status"],
        dados["telefone"],
        dados["nome"],
        dados["produto"],
        dados["cidade_uf"],
        dados["email"],
        dados["aplicacao"],
        dados["origem"],
        dados["consultor"],
        dados["observacao"],
        resumo,
        nome_whatsapp,
    ]

    for coluna, valor in enumerate(
        valores,
        start=1
    ):
        planilha.cell(
            row=linha,
            column=coluna,
            value=valor
        )

    planilha[
        f"F{linha}"
    ].number_format = "@"

    return (
        resumo,
        nome_whatsapp
    )


def criar_planilha():
    """
    Cria a planilha quando não existe.

    Quando já existe, verifica e realiza
    a migração da coluna UNIDADE.
    """

    preparar_ambiente()

    if ARQUIVO_LEADS.exists():
        arquivo, _ = (
            _abrir_planilha()
        )

        arquivo.close()

        return False

    _criar_planilha_nova()

    return True


def salvar_lead(lead):
    """
    Salva um novo lead.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        linha_vazia = (
            planilha.max_row + 1
        )

        dados = _montar_dados(
            lead
        )

        data_cadastro = (
            datetime.now().strftime(
                "%d/%m/%Y %H:%M"
            )
        )

        _escrever_linha(
            planilha,
            linha_vazia,
            data_cadastro,
            dados
        )

        _configurar_planilha(
            planilha
        )

        _salvar_seguro(
            arquivo
        )

        return linha_vazia

    except Exception:
        try:
            arquivo.close()

        except Exception:
            pass

        raise


def pegar_leads():
    """
    Retorna todos os leads cadastrados.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        leads = []

        for linha in range(
            2,
            planilha.max_row + 1
        ):
            telefone = (
                planilha[
                    f"F{linha}"
                ].value
            )

            nome = (
                planilha[
                    f"G{linha}"
                ].value
            )

            if (
                telefone is None
                and nome is None
            ):
                continue

            leads.append(
                _lead_da_linha(
                    planilha,
                    linha
                )
            )

        return leads

    finally:
        arquivo.close()


def listar_leads():
    """
    Mantida para compatibilidade.
    """

    return pegar_leads()


def buscar_lead_por_nome(
    nome_busca
):
    """
    Busca um lead pelo nome e retorna
    a linha da planilha.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        nome_procurado = (
            _valor(
                nome_busca
            ).casefold()
        )

        for linha in range(
            2,
            planilha.max_row + 1
        ):
            nome = _valor(
                planilha[
                    f"G{linha}"
                ].value
            )

            if (
                nome.casefold()
                == nome_procurado
            ):
                return linha

        return None

    finally:
        arquivo.close()


def buscar_lead_por_telefone(
    telefone_busca
):
    """
    Busca um lead pelo telefone,
    ignorando a formatação.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        telefone_procurado = (
            _normalizar_telefone(
                telefone_busca
            )
        )

        for linha in range(
            2,
            planilha.max_row + 1
        ):
            telefone_planilha = (
                _normalizar_telefone(
                    planilha[
                        f"F{linha}"
                    ].value
                )
            )

            if (
                telefone_planilha
                == telefone_procurado
            ):
                return _lead_da_linha(
                    planilha,
                    linha
                )

        return None

    finally:
        arquivo.close()


def atualizar_lead(lead):
    """
    Atualiza um lead pela chave linha.
    """

    if "linha" not in lead:
        raise KeyError(
            "O lead precisa possuir "
            "a chave 'linha'."
        )

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        linha = int(
            lead["linha"]
        )

        if (
            linha < 2
            or linha > planilha.max_row
        ):
            raise ValueError(
                "Linha do lead inválida."
            )

        atual = _lead_da_linha(
            planilha,
            linha
        )

        dados = _montar_dados(
            lead,
            atual
        )

        data_cadastro = (
            _valor(
                lead.get(
                    "data_cadastro",
                    lead.get(
                        "data",
                        atual[
                            "data_cadastro"
                        ]
                    )
                )
            )
            or atual[
                "data_cadastro"
            ]
        )

        resumo, whatsapp = (
            _escrever_linha(
                planilha,
                linha,
                data_cadastro,
                dados
            )
        )

        _configurar_planilha(
            planilha
        )

        _salvar_seguro(
            arquivo
        )

        dados[
            "data_cadastro"
        ] = data_cadastro

        dados["linha"] = linha
        dados["resumo"] = resumo
        dados["whatsapp"] = whatsapp

        return dados

    except Exception:
        try:
            arquivo.close()

        except Exception:
            pass

        raise


def excluir_lead(linha):
    """
    Exclui uma linha da planilha.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        linha = int(
            linha
        )

        if (
            linha < 2
            or linha > planilha.max_row
        ):
            raise ValueError(
                "Linha do lead inválida."
            )

        planilha.delete_rows(
            linha
        )

        _configurar_planilha(
            planilha
        )

        _salvar_seguro(
            arquivo
        )

        return True

    except Exception:
        try:
            arquivo.close()

        except Exception:
            pass

        raise


def atualizar_status(
    telefone,
    novo_status
):
    """
    Atualiza o status pelo telefone
    e registra a última interação.
    """

    arquivo, planilha = (
        _abrir_planilha()
    )

    try:
        telefone_procurado = (
            _normalizar_telefone(
                telefone
            )
        )

        status_novo = _valor(
            novo_status
        )

        for linha in range(
            2,
            planilha.max_row + 1
        ):
            telefone_planilha = (
                _normalizar_telefone(
                    planilha[
                        f"F{linha}"
                    ].value
                )
            )

            if (
                telefone_planilha
                != telefone_procurado
            ):
                continue

            planilha[
                f"E{linha}"
            ] = status_novo

            planilha[
                f"C{linha}"
            ] = (
                datetime.now().strftime(
                    "%d/%m/%Y %H:%M"
                )
            )

            if (
                status_novo
                .strip()
                .upper()
                == "DECLINADO"
            ):
                planilha[
                    f"B{linha}"
                ] = ""

            _configurar_planilha(
                planilha
            )

            _salvar_seguro(
                arquivo
            )

            return True

        arquivo.close()

        return False

    except Exception:
        try:
            arquivo.close()

        except Exception:
            pass

        raise


def dashboard():
    """
    Retorna as quantidades por status
    e a taxa de conversão.
    """

    leads = pegar_leads()

    contagem = {
        status: 0
        for status in STATUS_VALIDOS
    }

    for lead in leads:
        status = (
            _valor(
                lead.get(
                    "status"
                )
            )
            or "EM ANDAMENTO"
        )

        if status not in contagem:
            contagem[status] = 0

        contagem[status] += 1

    total = len(
        leads
    )

    conquistados = (
        contagem.get(
            "CONQUISTADO - SUPRIM",
            0
        )
        +
        contagem.get(
            "CONQUISTADO - EQUIP",
            0
        )
    )

    taxa_conversao = (
        conquistados / total * 100
        if total
        else 0
    )

    return {
        "total": total,
        "status": contagem,
        "taxa_conversao": taxa_conversao,
    }


def _converter_data_contato(valor):
    """
    Converte o próximo contato para data.
    """

    if isinstance(
        valor,
        datetime
    ):
        return valor.date()

    if isinstance(
        valor,
        date
    ):
        return valor

    texto = _valor(
        valor
    )

    if not texto:
        return None

    for formato in (
        "%d/%m/%Y",
        "%d/%m/%Y %H:%M",
    ):
        try:
            return datetime.strptime(
                texto,
                formato
            ).date()

        except ValueError:
            continue

    return None


def listar_followups():
    """
    Função mantida para compatibilidade.

    Leads declinados ficam em Finalizado.
    Leads ativos sem data não são exibidos.
    """

    hoje = datetime.now().date()

    grupos = {
        "atrasados": [],
        "hoje": [],
        "proximos": [],
        "finalizados": [],
        "invalidos": [],
    }

    for lead in pegar_leads():
        item = lead.copy()

        status = _valor(
            lead.get(
                "status",
                ""
            )
        ).upper()

        if status == "DECLINADO":
            grupos[
                "finalizados"
            ].append(
                item
            )
            continue

        texto_data = _valor(
            lead.get(
                "proximo_contato",
                ""
            )
        )

        # Leads ativos sem data ficam fora
        # do Follow-up.
        if not texto_data:
            continue

        data_contato = (
            _converter_data_contato(
                texto_data
            )
        )

        if data_contato is None:
            grupos[
                "invalidos"
            ].append(
                item
            )
            continue

        item[
            "_data_contato"
        ] = data_contato

        if data_contato < hoje:
            grupo = "atrasados"

        elif data_contato == hoje:
            grupo = "hoje"

        else:
            grupo = "proximos"

        grupos[
            grupo
        ].append(
            item
        )

    def ordenar_data(item):
        return (
            item["_data_contato"],
            _valor(
                item.get(
                    "nome",
                    ""
                )
            ).casefold()
        )

    for grupo in (
        "atrasados",
        "hoje",
        "proximos",
    ):
        grupos[grupo].sort(
            key=ordenar_data
        )

        for item in grupos[grupo]:
            item.pop(
                "_data_contato",
                None
            )

    for grupo in (
        "finalizados",
        "invalidos",
    ):
        grupos[grupo].sort(
            key=lambda item: _valor(
                item.get(
                    "nome",
                    ""
                )
            ).casefold()
        )

    return grupos


if __name__ == "__main__":
    criado = criar_planilha()

    if criado:
        print(
            "Planilha criada com sucesso."
        )

    else:
        print(
            "Planilha verificada com sucesso."
        )

    print(
        ARQUIVO_LEADS
    )