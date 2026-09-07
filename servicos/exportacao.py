from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)


COLUNAS_EXPORTACAO = [
    (
        "data_cadastro",
        "DATA DE CADASTRO"
    ),
    (
        "proximo_contato",
        "PRÓXIMO CONTATO"
    ),
    (
        "ultima_interacao",
        "ÚLTIMA INTERAÇÃO"
    ),
    (
        "unidade",
        "UNIDADE"
    ),
    (
        "status",
        "STATUS"
    ),
    (
        "telefone",
        "TELEFONE"
    ),
    (
        "nome",
        "NOME"
    ),
    (
        "produto",
        "PRODUTO"
    ),
    (
        "cidade_uf",
        "CIDADE / UF"
    ),
    (
        "email",
        "E-MAIL"
    ),
    (
        "aplicacao",
        "APLICAÇÃO"
    ),
    (
        "origem",
        "ORIGEM"
    ),
    (
        "consultor",
        "CONSULTOR"
    ),
    (
        "observacao",
        "OBSERVAÇÃO"
    ),
    (
        "resumo",
        "RESUMO"
    ),
    (
        "whatsapp",
        "NOME NO WHATSAPP"
    ),
]


LARGURAS = {
    "A": 20,
    "B": 20,
    "C": 20,
    "D": 20,
    "E": 25,
    "F": 18,
    "G": 25,
    "H": 25,
    "I": 20,
    "J": 30,
    "K": 25,
    "L": 20,
    "M": 20,
    "N": 55,
    "O": 45,
    "P": 55,
}


def exportar_leads(
    lista_leads,
    caminho
):
    """
    Exporta os leads recebidos para
    uma nova planilha do Excel.
    """

    if not lista_leads:
        raise ValueError(
            "Não existem leads para exportar."
        )

    destino = Path(
        caminho
    )

    if (
        destino.suffix.lower()
        != ".xlsx"
    ):
        destino = destino.with_suffix(
            ".xlsx"
        )

    workbook = Workbook()

    planilha = workbook.active
    planilha.title = (
        "Relatório de Leads"
    )

    cabecalhos = [
        titulo
        for _, titulo
        in COLUNAS_EXPORTACAO
    ]

    planilha.append(
        cabecalhos
    )

    preenchimento = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    for celula in planilha[1]:
        celula.font = Font(
            bold=True
        )

        celula.fill = (
            preenchimento
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    planilha.row_dimensions[
        1
    ].height = 30

    for lead in lista_leads:
        valores = [
            lead.get(
                chave,
                ""
            )
            for chave, _
            in COLUNAS_EXPORTACAO
        ]

        planilha.append(
            valores
        )

    planilha.freeze_panes = "A2"

    planilha.auto_filter.ref = (
        planilha.dimensions
    )

    for coluna, largura in (
        LARGURAS.items()
    ):
        planilha.column_dimensions[
            coluna
        ].width = largura

    for linha in planilha.iter_rows(
        min_row=2,
        max_col=16
    ):
        for celula in linha:
            celula.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # O telefone agora está na coluna F.
    for numero_linha in range(
        2,
        planilha.max_row + 1
    ):
        planilha[
            f"F{numero_linha}"
        ].number_format = "@"

    try:
        workbook.save(
            destino
        )

    except PermissionError as erro:
        raise PermissionError(
            "Não foi possível salvar o relatório.\n\n"
            "Feche o arquivo no Excel "
            "e tente novamente."
        ) from erro

    finally:
        workbook.close()

    return destino